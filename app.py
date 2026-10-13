import base64
import hashlib
import html as _html
import io
import json
import os
import re
import shutil
import sqlite3
import subprocess
from datetime import datetime
from pathlib import Path

import pandas as pd
import streamlit as st
import streamlit.components.v1 as components

# =====================================================================
# KONFIGURASI
# Default: folder yang sama dengan file ini. Kalau data lama kamu ada di
# /Users/macbook/SIN dan file ini ada di tempat lain, jalankan dengan:
#   SIN_DIR=/Users/macbook/SIN streamlit run app.py
# =====================================================================
BASE_DIR = Path(os.environ.get("SIN_DIR", Path(__file__).resolve().parent))
DATA_DIR = BASE_DIR / "data"
DOCS_DIR = DATA_DIR / "documents"
PREVIEW_DIR = DATA_DIR / "preview_cache"
DB_PATH = BASE_DIR / "sin_archive.db"

for _d in (BASE_DIR, DOCS_DIR, PREVIEW_DIR):
    _d.mkdir(parents=True, exist_ok=True)

STATUS_OPTIONS = ["Tidak Ada", "Ada - Asli", "Ada - Copy/Scan"]
ADA = ("Ada - Asli", "Ada - Copy/Scan")
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp"}
SHEET_EXT = {".xlsx", ".xlsm", ".xls", ".csv"}
OFFICE_EXT = {".docx", ".doc", ".odt", ".rtf", ".pptx", ".ppt"}

PAGE_SEARCH = "🔎 Pencarian"
PAGE_DEBTOR = "📂 Dokumen Debitur"
PAGE_ADD = "➕ Tambah Debitur"

CHECKLIST_STRUCTURE = {
    "1. Surat Permohonan Kredit": ["Surat Permohonan"],
    "2. Identitas Debitur": [
        "Copy bukti diri/KTP",
        "Copy surat kewarganegaraan/surat keterangan ganti nama",
        "Foto debitur",
        "Identitas debitur lainnya",
    ],
    "3. Identitas usaha debitur": [
        "Akte pendirian/perubahan", "SIUP", "SITU", "TDP", "SIUJK", "NPWP", "Lainnya",
    ],
    "4. Surat Hutang": ["Surat hutang", "Addendum surat hutang"],
    "5. Pengikatan Agunan": [
        "Sertifikat CV", "Sertifikat Hipotik", "Sertifikat HT", "Fiducia",
        "Gadai", "Cessie", "SKMHT", "Personal/Corporate Guarantee", "Lainnya",
    ],
    "6. Bukti Kepemilikan Agunan": [
        "Hak atas tanah", "Bilyet deposito/Buku tabungan", "BPKB",
        "Promissory notes/Commercial paper/Saham", "Bank Garansi", "Faktur/Kwitansi", "Lainnya",
    ],
    "7. Paket Kredit": [
        "MAK", "LKN", "Penilaian Agunan", "Pemeriksaan kelengkapan paket",
        "PTK", "Offering letter", "PPND", "IPK", "Surat lainnya",
    ],
    "8. Surat-surat lainnya": ["Asuransi", "Surat-surat lainnya"],
    "9. Pemeriksaan Agunan": ["Penilaian tanah", "Penilaian tanah dan bangunan", "Lainnya"],
    "10. Catatan Lain-lain": ["Catatan Lain-lain"],
}
TOTAL_ITEMS = sum(len(v) for v in CHECKLIST_STRUCTURE.values())
ITEM_LABELS = [f"{c} | {s}" for c, subs in CHECKLIST_STRUCTURE.items() for s in subs]
LABEL_TO_ITEM = {f"{c} | {s}": (c, s) for c, subs in CHECKLIST_STRUCTURE.items() for s in subs}
DEFAULT_LABEL = "10. Catatan Lain-lain | Catatan Lain-lain"

# Tebakan otomatis kategori dari nama file (token di nama file -> item checklist)
KEYWORDS = {
    "permohonan": ("1. Surat Permohonan Kredit", "Surat Permohonan"),
    "ktp": ("2. Identitas Debitur", "Copy bukti diri/KTP"),
    "ktpel": ("2. Identitas Debitur", "Copy bukti diri/KTP"),
    "foto": ("2. Identitas Debitur", "Foto debitur"),
    "akte": ("3. Identitas usaha debitur", "Akte pendirian/perubahan"),
    "akta": ("3. Identitas usaha debitur", "Akte pendirian/perubahan"),
    "siup": ("3. Identitas usaha debitur", "SIUP"),
    "situ": ("3. Identitas usaha debitur", "SITU"),
    "tdp": ("3. Identitas usaha debitur", "TDP"),
    "siujk": ("3. Identitas usaha debitur", "SIUJK"),
    "npwp": ("3. Identitas usaha debitur", "NPWP"),
    "hutang": ("4. Surat Hutang", "Surat hutang"),
    "addendum": ("4. Surat Hutang", "Addendum surat hutang"),
    "sht": ("5. Pengikatan Agunan", "Sertifikat HT"),
    "apht": ("5. Pengikatan Agunan", "Sertifikat HT"),
    "ht": ("5. Pengikatan Agunan", "Sertifikat HT"),
    "fidusia": ("5. Pengikatan Agunan", "Fiducia"),
    "fiducia": ("5. Pengikatan Agunan", "Fiducia"),
    "gadai": ("5. Pengikatan Agunan", "Gadai"),
    "cessie": ("5. Pengikatan Agunan", "Cessie"),
    "skmht": ("5. Pengikatan Agunan", "SKMHT"),
    "guarantee": ("5. Pengikatan Agunan", "Personal/Corporate Guarantee"),
    "shm": ("6. Bukti Kepemilikan Agunan", "Hak atas tanah"),
    "sertifikat": ("6. Bukti Kepemilikan Agunan", "Hak atas tanah"),
    "bpkb": ("6. Bukti Kepemilikan Agunan", "BPKB"),
    "bilyet": ("6. Bukti Kepemilikan Agunan", "Bilyet deposito/Buku tabungan"),
    "deposito": ("6. Bukti Kepemilikan Agunan", "Bilyet deposito/Buku tabungan"),
    "tabungan": ("6. Bukti Kepemilikan Agunan", "Bilyet deposito/Buku tabungan"),
    "garansi": ("6. Bukti Kepemilikan Agunan", "Bank Garansi"),
    "faktur": ("6. Bukti Kepemilikan Agunan", "Faktur/Kwitansi"),
    "kwitansi": ("6. Bukti Kepemilikan Agunan", "Faktur/Kwitansi"),
    "mak": ("7. Paket Kredit", "MAK"),
    "lkn": ("7. Paket Kredit", "LKN"),
    "penilaian": ("7. Paket Kredit", "Penilaian Agunan"),
    "ptk": ("7. Paket Kredit", "PTK"),
    "offering": ("7. Paket Kredit", "Offering letter"),
    "ppnd": ("7. Paket Kredit", "PPND"),
    "ipk": ("7. Paket Kredit", "IPK"),
    "asuransi": ("8. Surat-surat lainnya", "Asuransi"),
}


# =====================================================================
# DATABASE (skema sama dengan versi lama, data lama tetap terbaca)
# =====================================================================
def get_connection():
    conn = sqlite3.connect(DB_PATH)
    conn.execute("PRAGMA foreign_keys = ON")
    return conn


def init_db():
    conn = get_connection()
    c = conn.cursor()
    c.execute("""
        CREATE TABLE IF NOT EXISTS debtors (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            nama_debitur TEXT NOT NULL,
            nomor_rekening TEXT,
            plafond REAL,
            rm_kredit TEXT,
            brimen TEXT,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        )""")
    c.execute("""
        CREATE TABLE IF NOT EXISTS debtor_bank_accounts (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            debtor_id INTEGER,
            nomor_rekening TEXT NOT NULL,
            nama_bank TEXT DEFAULT 'BRI',
            keterangan TEXT,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY(debtor_id) REFERENCES debtors(id) ON DELETE CASCADE
        )""")
    c.execute("""
        CREATE TABLE IF NOT EXISTS documents (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            debtor_id INTEGER,
            category TEXT,
            sub_item TEXT,
            status TEXT,
            perihal TEXT,
            keterangan TEXT,
            file_path TEXT,
            updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY(debtor_id) REFERENCES debtors(id) ON DELETE CASCADE
        )""")
    c.execute("CREATE INDEX IF NOT EXISTS idx_docs_debtor ON documents(debtor_id)")
    c.execute("CREATE INDEX IF NOT EXISTS idx_acc_debtor ON debtor_bank_accounts(debtor_id)")
    conn.commit()
    conn.close()


init_db()


# =====================================================================
# HELPER
# =====================================================================
def parse_file_paths(s):
    if not s:
        return []
    try:
        data = json.loads(s)
        return data if isinstance(data, list) else [s]
    except Exception:
        return [s]


def dump_file_paths(paths):
    return json.dumps(paths)


def is_present(status, files):
    return status in ADA or len(files) > 0


def safe_name(s):
    return re.sub(r"[^\w.\-]+", "_", s, flags=re.UNICODE).strip("._") or "file"


def save_upload(debtor_id, sub, uploaded):
    folder = DOCS_DIR / str(debtor_id)
    folder.mkdir(parents=True, exist_ok=True)
    name = f"{safe_name(sub)}_{safe_name(uploaded.name)}"
    dest = folder / name
    n = 1
    while dest.exists():
        dest = folder / f"{Path(name).stem}_{n}{Path(name).suffix}"
        n += 1
    dest.write_bytes(uploaded.getbuffer())
    return str(dest)


def guess_label(filename):
    tokens = set(re.split(r"[^a-z0-9]+", Path(filename).stem.lower()))
    for kw, (cat, sub) in KEYWORDS.items():
        if kw in tokens:
            return f"{cat} | {sub}"
    return DEFAULT_LABEL


def find_duplicate_rek(rek, exclude_id=-1):
    conn = get_connection()
    row = conn.execute("""
        SELECT d.nama_debitur FROM debtors d
        LEFT JOIN debtor_bank_accounts a ON a.debtor_id = d.id
        WHERE (d.nomor_rekening = ? OR a.nomor_rekening = ?) AND d.id != ?
        LIMIT 1
    """, (rek, rek, exclude_id)).fetchone()
    conn.close()
    return row[0] if row else None


def load_docs(debtor_id):
    """Return dict {(category, sub_item): {...}}"""
    conn = get_connection()
    df = pd.read_sql("SELECT * FROM documents WHERE debtor_id = ?", conn, params=(debtor_id,))
    conn.close()
    out = {}
    for _, r in df.iterrows():
        out[(r["category"], r["sub_item"])] = {
            "id": int(r["id"]),
            "status": r["status"] or "Tidak Ada",
            "perihal": r["perihal"] or "",
            "keterangan": r["keterangan"] or "",
            "files": parse_file_paths(r["file_path"]),
        }
    return out


def upsert_doc(conn, debtor_id, cat, sub, status, perihal, ket, files=None):
    row = conn.execute(
        "SELECT id, file_path FROM documents WHERE debtor_id=? AND category=? AND sub_item=?",
        (debtor_id, cat, sub),
    ).fetchone()
    if row:
        fp = row[1] if files is None else dump_file_paths(files)
        conn.execute("""
            UPDATE documents SET status=?, perihal=?, keterangan=?, file_path=?,
            updated_at=CURRENT_TIMESTAMP WHERE id=?
        """, (status, perihal, ket, fp, row[0]))
    else:
        conn.execute("""
            INSERT INTO documents (debtor_id, category, sub_item, status, perihal, keterangan, file_path)
            VALUES (?, ?, ?, ?, ?, ?, ?)
        """, (debtor_id, cat, sub, status, perihal, ket, dump_file_paths(files or [])))


def load_summary():
    conn = get_connection()
    debtors = pd.read_sql(
        "SELECT id, nama_debitur, nomor_rekening, plafond, rm_kredit, brimen "
        "FROM debtors ORDER BY nama_debitur COLLATE NOCASE", conn)
    accs = pd.read_sql("SELECT debtor_id, nomor_rekening FROM debtor_bank_accounts", conn)
    docs = pd.read_sql(
        "SELECT debtor_id, sub_item, status, perihal, keterangan, file_path FROM documents", conn)
    conn.close()

    acc_map = accs.groupby("debtor_id")["nomor_rekening"].apply(list).to_dict() if not accs.empty else {}
    doc_groups = {k: g for k, g in docs.groupby("debtor_id")} if not docs.empty else {}

    rows = []
    for _, d in debtors.iterrows():
        did = int(d["id"])
        reks = [d["nomor_rekening"]] + acc_map.get(did, [])
        reks = list(dict.fromkeys([r for r in reks if r]))
        done = 0
        hay = [d["nama_debitur"] or "", d["rm_kredit"] or "", d["brimen"] or ""] + reks
        g = doc_groups.get(did)
        if g is not None:
            for _, r in g.iterrows():
                files = parse_file_paths(r["file_path"])
                if is_present(r["status"], files):
                    done += 1
                hay += [r["sub_item"] or "", r["perihal"] or "", r["keterangan"] or ""]
                hay += [Path(f).name for f in files]
        done = min(done, TOTAL_ITEMS)
        rows.append({
            "id": did,
            "nama_debitur": d["nama_debitur"],
            "rekening": ", ".join(reks) if reks else "-",
            "plafond": float(d["plafond"] or 0),
            "rm_kredit": d["rm_kredit"] or "",
            "brimen": d["brimen"] or "",
            "done": done,
            "pct": round(done / TOTAL_ITEMS * 100, 1),
            "haystack": " ".join(hay).lower(),
        })
    return pd.DataFrame(rows, columns=[
        "id", "nama_debitur", "rekening", "plafond", "rm_kredit", "brimen", "done", "pct", "haystack"])


def open_debtor(debtor_id, flash=None):
    st.session_state["cur_debtor"] = int(debtor_id)
    st.session_state["preview_file"] = None
    st.session_state.pop("debtor_picker", None)
    st.session_state["nav_target"] = PAGE_DEBTOR
    if flash:
        st.session_state["flash"] = flash
    st.rerun()


# =====================================================================
# PREVIEW
# =====================================================================
def libreoffice_bin():
    found = shutil.which("soffice")
    if found:
        return found
    mac = "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    return mac if os.path.exists(mac) else None


def convert_to_pdf(path: Path):
    """Konversi Word/PPT ke PDF via LibreOffice (di-cache). None kalau tidak tersedia."""
    key = hashlib.sha1(f"{path}{path.stat().st_mtime}".encode()).hexdigest()[:16]
    out_dir = PREVIEW_DIR / key
    target = out_dir / f"{path.stem}.pdf"
    if target.exists():
        return target
    soffice = libreoffice_bin()
    if not soffice:
        return None
    out_dir.mkdir(parents=True, exist_ok=True)
    try:
        subprocess.run(
            [soffice, "--headless", "--convert-to", "pdf", "--outdir", str(out_dir), str(path)],
            capture_output=True, timeout=120, check=False)
    except Exception:
        return None
    return target if target.exists() else None


def render_pdf(path: Path):
    data = path.read_bytes()
    if hasattr(st, "pdf"):
        try:
            st.pdf(data, height=700)
            return
        except Exception:
            pass  # fallback ke iframe (butuh: pip install "streamlit[pdf]" untuk st.pdf)
    if len(data) > 20 * 1024 * 1024:
        st.warning("PDF terlalu besar untuk preview. Gunakan tombol Download.")
        return
    b64 = base64.b64encode(data).decode()
    st.markdown(
        f'<iframe src="data:application/pdf;base64,{b64}" width="100%" height="700"></iframe>',
        unsafe_allow_html=True)


def render_sheet(path: Path):
    try:
        if path.suffix.lower() == ".csv":
            try:
                df = pd.read_csv(path, header=None, nrows=1000, dtype=str)
            except UnicodeDecodeError:
                df = pd.read_csv(path, header=None, nrows=1000, dtype=str, encoding="latin-1")
        else:
            xf = pd.ExcelFile(path)
            sheet = xf.sheet_names[0]
            if len(xf.sheet_names) > 1:
                sheet = st.selectbox("Sheet", xf.sheet_names, key=f"sheet_{path.name}")
            df = pd.read_excel(xf, sheet_name=sheet, header=None, nrows=1000, dtype=str)
        df = df.dropna(how="all").dropna(axis=1, how="all").fillna("")
        st.dataframe(df, hide_index=True)
        st.caption("Menampilkan maks. 1000 baris pertama. Download file untuk melihat lengkap.")
    except Exception as e:
        st.error(f"Gagal membaca spreadsheet: {e}")


def render_docx_text(path: Path):
    try:
        from docx import Document
    except ImportError:
        st.info("Install `python-docx` untuk preview teks Word: `pip install python-docx`.")
        return
    try:
        doc = Document(path)
        text = "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())
        if text:
            st.markdown(text)
        for t in doc.tables:
            rows = [[c.text for c in r.cells] for r in t.rows]
            if rows:
                st.dataframe(pd.DataFrame(rows), hide_index=True)
        st.caption("Preview teks saja (tanpa tata letak). Pasang LibreOffice untuk tampilan seperti aslinya.")
    except Exception as e:
        st.error(f"Gagal membaca dokumen Word: {e}")


def render_preview(path_str):
    p = Path(path_str)
    ext = p.suffix.lower()
    if not p.exists():
        st.error("File tidak ditemukan di disk.")
        return
    if ext in IMAGE_EXT:
        st.image(str(p))
    elif ext == ".pdf":
        render_pdf(p)
    elif ext in SHEET_EXT:
        render_sheet(p)
    elif ext in OFFICE_EXT:
        with st.spinner("Menyiapkan preview..."):
            pdf = convert_to_pdf(p)
        if pdf:
            render_pdf(pdf)
        elif ext == ".docx":
            render_docx_text(p)
        else:
            st.info("Preview format ini butuh LibreOffice. Gunakan tombol Download.")
    elif ext == ".txt":
        st.text(p.read_text(errors="ignore")[:20000])
    else:
        st.info("Format ini belum bisa dipreview. Gunakan tombol Download.")


# =====================================================================
# LAPORAN KELENGKAPAN (HTML & EXCEL)
# =====================================================================
STATUS_STYLE = {
    # status: (warna latar, warna teks, label)
    "Ada - Asli": ("#d4edda", "#155724", "✔ Ada - Asli"),
    "Ada - Copy/Scan": ("#d1ecf1", "#0c5460", "✔ Ada - Copy/Scan"),
    "Tidak Ada": ("#f8d7da", "#721c24", "✘ Belum ada"),
}
_TD = "border:1px solid #999999;padding:6px;vertical-align:top;font-size:13px;"
_TH = ("border:1px solid #999999;padding:6px;font-size:13px;background:#343a40;"
       "color:#ffffff;text-align:left;")


def checklist_rows(docs):
    rows = []
    for cat, subs in CHECKLIST_STRUCTURE.items():
        for i, sub in enumerate(subs, 1):
            d = docs.get((cat, sub))
            files = [f for f in (d["files"] if d else []) if os.path.exists(f)]
            status = d["status"] if d else "Tidak Ada"
            if status == "Tidak Ada" and files:
                status = "Ada - Copy/Scan"
            rows.append({
                "Kategori": cat, "No": i, "Item": sub, "Status": status,
                "Ada": status in ADA,
                "Perihal": d["perihal"] if d else "",
                "Keterangan": d["keterangan"] if d else "",
                "Jumlah File": len(files),
            })
    return rows


def report_summary(rows):
    total = len(rows)
    asli = sum(r["Status"] == "Ada - Asli" for r in rows)
    copy = sum(r["Status"] == "Ada - Copy/Scan" for r in rows)
    ada = asli + copy
    return {"total": total, "asli": asli, "copy": copy, "ada": ada,
            "belum": total - ada, "pct": (ada / total * 100) if total else 0}


def report_body_html(info, reks, rows, only_missing=False):
    """HTML dengan inline style (tanpa <style>), supaya bentuk tabel tetap
    utuh saat di-copy-paste ke Word, Outlook, Google Docs, dsb."""
    e = _html.escape
    s = report_summary(rows)
    show = [r for r in rows if not (only_missing and r["Ada"])]
    now = datetime.now().strftime("%d-%m-%Y %H:%M")

    meta = [
        ("Nama Debitur", info["nama_debitur"]),
        ("Nomor Rekening", reks),
        ("Plafond", f"Rp {float(info['plafond'] or 0):,.0f}"),
        ("RM Kredit", info["rm_kredit"] or "-"),
        ("BRIMEN", info["brimen"] or "-"),
        ("Tanggal Laporan", now),
    ]
    meta_html = "".join(
        f'<tr><td style="{_TD}width:160px;font-weight:bold;background:#f8f9fa;">{e(k)}</td>'
        f'<td style="{_TD}">{e(str(v))}</td></tr>' for k, v in meta)

    summary = (
        f'<p style="font-size:14px;margin:12px 0;"><b>Kelengkapan: {s["ada"]} dari {s["total"]} '
        f'item ({s["pct"]:.0f}%)</b> &nbsp;|&nbsp; Asli: {s["asli"]} &nbsp;|&nbsp; '
        f'Copy/Scan: {s["copy"]} &nbsp;|&nbsp; Belum ada: {s["belum"]}</p>')

    head = ("<tr>"
            f'<th style="{_TH}width:40px;">No</th>'
            f'<th style="{_TH}">Item Dokumen</th>'
            f'<th style="{_TH}width:140px;">Status</th>'
            f'<th style="{_TH}">Perihal</th>'
            f'<th style="{_TH}">Keterangan</th>'
            f'<th style="{_TH}width:50px;">File</th>'
            "</tr>")

    parts, last_cat = [], None
    for r in show:
        if r["Kategori"] != last_cat:
            last_cat = r["Kategori"]
            parts.append(
                f'<tr><td colspan="6" style="{_TD}background:#e9ecef;font-weight:bold;">'
                f'{e(last_cat)}</td></tr>')
        bg, fg, label = STATUS_STYLE[r["Status"]]
        parts.append(
            "<tr>"
            f'<td style="{_TD}text-align:center;">{r["No"]}</td>'
            f'<td style="{_TD}">{e(r["Item"])}</td>'
            f'<td style="{_TD}background:{bg};color:{fg};font-weight:bold;">{label}</td>'
            f'<td style="{_TD}">{e(r["Perihal"])}</td>'
            f'<td style="{_TD}">{e(r["Keterangan"])}</td>'
            f'<td style="{_TD}text-align:center;">{r["Jumlah File"]}</td>'
            "</tr>")
    if not parts:
        parts.append(f'<tr><td colspan="6" style="{_TD}text-align:center;">'
                     "Semua item dokumen sudah lengkap.</td></tr>")

    return (
        '<div id="lap" style="font-family:Arial,Helvetica,sans-serif;color:#212529;'
        'background:#ffffff;padding:12px;">'
        '<h2 style="margin:0 0 10px 0;">Laporan Kelengkapan Dokumen Kredit</h2>'
        '<table border="1" cellspacing="0" cellpadding="0" '
        f'style="border-collapse:collapse;width:100%;max-width:640px;">{meta_html}</table>'
        f"{summary}"
        '<table border="1" cellspacing="0" cellpadding="0" '
        f'style="border-collapse:collapse;width:100%;">{head}{"".join(parts)}</table>'
        "</div>")


def report_full_html(body):
    return ('<!DOCTYPE html><html lang="id"><head><meta charset="utf-8">'
            "<title>Laporan Kelengkapan Dokumen</title></head><body>"
            f"{body}</body></html>")


def build_report_xlsx(info, reks, rows, only_missing=False):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    s = report_summary(rows)
    show = [r for r in rows if not (only_missing and r["Ada"])]
    wb = Workbook()
    ws = wb.active
    ws.title = "Kelengkapan"
    bold = Font(bold=True)
    thin = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws["A1"] = "Laporan Kelengkapan Dokumen Kredit"
    ws["A1"].font = Font(bold=True, size=14)
    meta = [
        ("Nama Debitur", info["nama_debitur"]),
        ("Nomor Rekening", reks),
        ("Plafond (Rp)", float(info["plafond"] or 0)),
        ("RM Kredit", info["rm_kredit"] or "-"),
        ("BRIMEN", info["brimen"] or "-"),
        ("Kelengkapan", f"{s['ada']} / {s['total']} ({s['pct']:.0f}%)"),
        ("Tanggal Laporan", datetime.now().strftime("%d-%m-%Y %H:%M")),
    ]
    for i, (k, v) in enumerate(meta, start=3):
        ws.cell(i, 1, k).font = bold
        c = ws.cell(i, 2, v)
        c.alignment = Alignment(horizontal="left")
        if k.startswith("Plafond"):
            c.number_format = "#,##0"

    header_row = 3 + len(meta) + 1
    headers = ["Kategori", "No", "Item", "Status", "Perihal", "Keterangan", "Jumlah File"]
    for j, h in enumerate(headers, 1):
        c = ws.cell(header_row, j, h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="343A40")
        c.border = border

    fills = {"Ada - Asli": "D4EDDA", "Ada - Copy/Scan": "D1ECF1", "Tidak Ada": "F8D7DA"}
    r_idx = header_row
    for r in show:
        r_idx += 1
        values = [r["Kategori"], r["No"], r["Item"], STATUS_STYLE[r["Status"]][2],
                  r["Perihal"], r["Keterangan"], r["Jumlah File"]]
        for j, v in enumerate(values, 1):
            c = ws.cell(r_idx, j, v)
            c.border = border
            c.alignment = Alignment(vertical="top", wrap_text=True)
        ws.cell(r_idx, 4).fill = PatternFill("solid", fgColor=fills[r["Status"]])
        ws.cell(r_idx, 4).font = bold

    ws.auto_filter.ref = f"A{header_row}:G{max(r_idx, header_row + 1)}"
    ws.freeze_panes = ws.cell(header_row + 1, 1)
    for col, w in zip("ABCDEFG", [30, 6, 42, 20, 28, 28, 12]):
        ws.column_dimensions[col].width = w

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def report_section(debtor_id, info, reks, docs):
    st.divider()
    st.subheader("📊 Laporan Kelengkapan")
    st.caption("Tabel kelengkapan dokumen siap pakai untuk laporan ke atasan. "
               "HTML cocok untuk di-copy-paste ke Word/email, Excel untuk diolah lagi.")
    rows = checklist_rows(docs)
    s = report_summary(rows)
    m1, m2, m3, m4 = st.columns(4)
    m1.metric("Lengkap", f"{s['ada']} / {s['total']}")
    m2.metric("Asli", s["asli"])
    m3.metric("Copy/Scan", s["copy"])
    m4.metric("Belum ada", s["belum"])

    only_missing = st.checkbox("Hanya tampilkan yang belum lengkap",
                               key=f"rep_missing_{debtor_id}")
    body = report_body_html(info, reks, rows, only_missing)
    fname = f"Kelengkapan_{safe_name(info['nama_debitur'])}_{datetime.now():%Y%m%d}"

    d1, d2, _ = st.columns([1, 1, 3])
    d1.download_button("📥 Download HTML", report_full_html(body).encode("utf-8"),
                       file_name=f"{fname}.html", mime="text/html",
                       key=f"rep_html_{debtor_id}")
    d2.download_button("📥 Download Excel", build_report_xlsx(info, reks, rows, only_missing),
                       file_name=f"{fname}.xlsx",
                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                       key=f"rep_xlsx_{debtor_id}")

    copy_btn = (
        '<button style="margin:0 0 8px 0;padding:6px 12px;cursor:pointer;" '
        "onclick=\"var r=document.createRange();r.selectNode(document.getElementById('lap'));"
        "var s=window.getSelection();s.removeAllRanges();s.addRange(r);"
        "document.execCommand('copy');this.innerText='✅ Tersalin';\">📋 Salin tabel</button>")
    with st.expander("👁️ Pratinjau tabel", expanded=True):
        components.html(copy_btn + body, height=620, scrolling=True)


# =====================================================================
# HALAMAN 1: PENCARIAN
# =====================================================================
def page_search():
    st.header("🔎 Cari Debitur & Dokumen")
    df = load_summary()
    if df.empty:
        st.info("Belum ada debitur. Mulai dari menu **➕ Tambah Debitur**.")
        return

    c1, c2, c3 = st.columns([3, 1, 1])
    q = c1.text_input(
        "Cari", placeholder="Nama, nomor rekening, perihal, atau nama file...").strip().lower()
    rm_opts = ["Semua"] + sorted({x for x in df["rm_kredit"] if x})
    br_opts = ["Semua"] + sorted({x for x in df["brimen"] if x})
    rm = c2.selectbox("RM Kredit", rm_opts)
    br = c3.selectbox("BRIMEN", br_opts)

    if q:
        df = df[df["haystack"].str.contains(q, regex=False)]
    if rm != "Semua":
        df = df[df["rm_kredit"] == rm]
    if br != "Semua":
        df = df[df["brimen"] == br]
    df = df.reset_index(drop=True)

    m1, m2, m3 = st.columns(3)
    m1.metric("Debitur ditemukan", len(df))
    m2.metric("Total Plafond (Rp)", f"{df['plafond'].sum():,.0f}")
    m3.metric("Rata-rata kelengkapan", f"{df['pct'].mean():.0f}%" if len(df) else "-")

    if df.empty:
        st.warning("Tidak ada hasil yang cocok.")
        return

    st.caption("Klik satu baris untuk membuka dokumen debitur.")
    view = df[["nama_debitur", "rekening", "plafond", "rm_kredit", "brimen", "pct"]]
    event = st.dataframe(
        view,
        hide_index=True,
        on_select="rerun",
        selection_mode="single-row",
        key="debtor_table",
        column_config={
            "nama_debitur": "Nama Debitur",
            "rekening": "Rekening",
            "plafond": st.column_config.NumberColumn("Plafond (Rp)", format="%d"),
            "rm_kredit": "RM Kredit",
            "brimen": "BRIMEN",
            "pct": st.column_config.ProgressColumn(
                "Kelengkapan", min_value=0, max_value=100, format="%.0f%%"),
        },
    )
    rows = event.selection.rows
    if rows:
        open_debtor(df.iloc[rows[0]]["id"])


# =====================================================================
# HALAMAN 2: TAMBAH DEBITUR (profil saja)
# =====================================================================
def page_add():
    st.header("➕ Tambah Debitur Baru")
    st.caption("Cukup isi profil. Dokumen diunggah setelah ini di halaman dokumen debitur.")
    k = st.session_state.setdefault("form_reset_key", 0)

    with st.form(f"add_form_{k}"):
        nama = st.text_input("Nama Debitur", key=f"a_nama_{k}")
        c1, c2 = st.columns(2)
        rek = c1.text_input("Nomor Rekening Utama", key=f"a_rek_{k}")
        extra = c2.text_input("Rekening Tambahan (pisahkan koma)", key=f"a_extra_{k}")
        plafond = st.number_input("Plafond (Rp)", min_value=0.0, step=1_000_000.0,
                                  format="%.0f", key=f"a_plafond_{k}")
        c3, c4 = st.columns(2)
        rm = c3.text_input("RM Kredit", key=f"a_rm_{k}")
        brimen = c4.text_input("BRIMEN", key=f"a_brimen_{k}")
        submitted = st.form_submit_button("Simpan & Lanjut ke Dokumen", type="primary")

    if not submitted:
        return

    nama, rek = nama.strip(), rek.strip()
    if not nama or not rek:
        st.error("Nama Debitur dan Nomor Rekening Utama wajib diisi.")
        return
    extras = [r.strip() for r in extra.split(",") if r.strip()]
    all_reks = list(dict.fromkeys([rek] + extras))
    for r in all_reks:
        owner = find_duplicate_rek(r)
        if owner:
            st.error(f"Rekening {r} sudah terdaftar atas nama **{owner}**.")
            return

    saved_id = None
    conn = get_connection()
    try:
        cur = conn.execute(
            "INSERT INTO debtors (nama_debitur, nomor_rekening, plafond, rm_kredit, brimen) "
            "VALUES (?, ?, ?, ?, ?)", (nama, rek, plafond, rm.strip(), brimen.strip()))
        saved_id = cur.lastrowid
        for r in all_reks[1:]:
            conn.execute(
                "INSERT INTO debtor_bank_accounts (debtor_id, nomor_rekening) VALUES (?, ?)",
                (saved_id, r))
        conn.commit()
        (DOCS_DIR / str(saved_id)).mkdir(parents=True, exist_ok=True)
    except Exception as e:
        conn.rollback()
        saved_id = None
        st.error(f"Gagal menyimpan data: {e}")
    finally:
        conn.close()

    if saved_id:
        st.session_state["form_reset_key"] += 1
        open_debtor(saved_id, flash=f"Debitur {nama} berhasil disimpan. Silakan unggah dokumen.")


# =====================================================================
# HALAMAN 3: DOKUMEN DEBITUR
# =====================================================================
def tab_documents(debtor_id, docs):
    # indeks semua file yang ada di disk
    file_index = {}
    for (cat, sub), d in docs.items():
        for f in d["files"]:
            if os.path.exists(f):
                file_index[f] = {"cat": cat, "sub": sub, "doc": d}

    sel = st.session_state.get("preview_file")
    if sel not in file_index:
        sel = next(iter(file_index), None)
        st.session_state["preview_file"] = sel

    left, right = st.columns([2, 3])

    with left:
        flt = st.text_input("Filter dokumen", placeholder="nama item / perihal / file",
                            key=f"flt_{debtor_id}").strip().lower()
        show_missing = st.toggle("Tampilkan yang belum ada", key=f"miss_{debtor_id}")
        box = st.container(height=640)
        with box:
            any_shown = False
            for cat, subs in CHECKLIST_STRUCTURE.items():
                block = []
                for sub in subs:
                    d = docs.get((cat, sub))
                    files = [f for f in (d["files"] if d else []) if f in file_index]
                    present = bool(d) and is_present(d["status"], files)
                    if flt:
                        text = " ".join([sub, d["perihal"], d["keterangan"]] +
                                        [Path(f).name for f in files]).lower() if d else sub.lower()
                        if flt not in text:
                            continue
                    if present or show_missing:
                        block.append((sub, d, files, present))
                if not block:
                    continue
                any_shown = True
                st.markdown(f"**{cat}**")
                for sub, d, files, present in block:
                    if files:
                        for i, f in enumerate(files):
                            label = f"📄 {sub}" + (f" ({i + 1})" if len(files) > 1 else "")
                            if st.button(label, key=f"pick_{hashlib.md5(f.encode()).hexdigest()}",
                                         type="primary" if f == sel else "secondary",
                                         help=Path(f).name):
                                st.session_state["preview_file"] = f
                                sel = f
                            if d["perihal"]:
                                st.caption(d["perihal"])
                    elif present:
                        st.caption(f"🟢 {sub}: {d['status']} (belum ada scan)")
                    else:
                        st.caption(f"🔴 {sub}: belum ada")
            if not any_shown:
                st.caption("Belum ada dokumen. Unggah lewat tab **⬆️ Upload**.")

    with right:
        if not sel:
            st.info("Pilih dokumen di sebelah kiri untuk melihat preview.")
            return
        meta = file_index[sel]
        st.markdown(f"**{meta['sub']}** · {Path(sel).name}")
        st.caption(f"{meta['cat']}"
                   + (f" · {meta['doc']['perihal']}" if meta['doc']['perihal'] else "")
                   + (f" · {meta['doc']['keterangan']}" if meta['doc']['keterangan'] else ""))
        b1, b2, _ = st.columns([1, 1, 3])
        with b1:
            st.download_button("📥 Download", data=Path(sel).read_bytes(),
                               file_name=Path(sel).name, key=f"dl_{debtor_id}")
        with b2:
            with st.popover("🗑️ Hapus"):
                st.warning("File akan dihapus permanen.")
                if st.button("Ya, hapus", key=f"delf_{debtor_id}", type="primary"):
                    try:
                        os.remove(sel)
                    except FileNotFoundError:
                        pass
                    d = meta["doc"]
                    remaining = [f for f in d["files"] if f != sel]
                    status = d["status"]
                    if not remaining and status != "Ada - Asli":
                        status = "Tidak Ada"
                    conn = get_connection()
                    conn.execute("UPDATE documents SET file_path=?, status=?, "
                                 "updated_at=CURRENT_TIMESTAMP WHERE id=?",
                                 (dump_file_paths(remaining), status, d["id"]))
                    conn.commit()
                    conn.close()
                    st.session_state["preview_file"] = None
                    st.session_state["flash"] = "File dihapus."
                    st.rerun()
        render_preview(sel)


def tab_upload(debtor_id, docs):
    st.caption("Pilih banyak file sekaligus. Kategori ditebak dari nama file "
               "(misal `KTP_budi.pdf`), cek dan koreksi di tabel sebelum disimpan.")
    bk = st.session_state.setdefault("bulk_key", 0)
    files = st.file_uploader(
        "File (PDF, gambar, Word, Excel)", accept_multiple_files=True,
        key=f"bulk_{debtor_id}_{bk}")
    if not files:
        return
    sig = hashlib.md5("|".join(f"{f.name}{f.size}" for f in files).encode()).hexdigest()[:8]
    base = pd.DataFrame({
        "File": [f.name for f in files],
        "Item Checklist": [guess_label(f.name) for f in files],
        "Status": ["Ada - Copy/Scan"] * len(files),
        "Perihal": [""] * len(files),
    })
    edited = st.data_editor(
        base, hide_index=True, num_rows="fixed", disabled=["File"],
        key=f"bulk_tbl_{debtor_id}_{sig}",
        column_config={
            "Item Checklist": st.column_config.SelectboxColumn(
                options=ITEM_LABELS, required=True, width="large"),
            "Status": st.column_config.SelectboxColumn(
                options=["Ada - Asli", "Ada - Copy/Scan"], required=True),
        })
    if st.button(f"💾 Simpan {len(files)} file", type="primary", key=f"bulk_save_{debtor_id}"):
        conn = get_connection()
        try:
            cache = {k: list(v["files"]) for k, v in docs.items()}
            for f, (_, row) in zip(files, edited.iterrows()):
                cat, sub = LABEL_TO_ITEM[row["Item Checklist"]]
                path = save_upload(debtor_id, sub, f)
                cache.setdefault((cat, sub), []).append(path)
                existing = docs.get((cat, sub), {})
                upsert_doc(conn, debtor_id, cat, sub, row["Status"],
                           row["Perihal"] or existing.get("perihal", ""),
                           existing.get("keterangan", ""), cache[(cat, sub)])
            conn.commit()
            st.session_state["bulk_key"] += 1
            st.session_state["flash"] = f"{len(files)} file berhasil disimpan."
        except Exception as e:
            conn.rollback()
            st.error(f"Gagal menyimpan: {e}")
            conn.close()
            return
        conn.close()
        st.rerun()


def tab_checklist(debtor_id, docs, info, reks):
    st.caption("Atur status, perihal, dan keterangan per kategori. Satu tombol simpan per kategori.")
    cat = st.selectbox("Kategori", list(CHECKLIST_STRUCTURE.keys()), key=f"chk_cat_{debtor_id}")
    rows = []
    for sub in CHECKLIST_STRUCTURE[cat]:
        d = docs.get((cat, sub))
        files = [f for f in (d["files"] if d else []) if os.path.exists(f)]
        status = d["status"] if d else "Tidak Ada"
        if status == "Tidak Ada" and files:
            status = "Ada - Copy/Scan"
        rows.append({"Item": sub, "Status": status,
                     "Perihal": d["perihal"] if d else "",
                     "Keterangan": d["keterangan"] if d else "",
                     "File": len(files)})
    edited = st.data_editor(
        pd.DataFrame(rows), hide_index=True, num_rows="fixed",
        disabled=["Item", "File"], key=f"chk_{debtor_id}_{cat}",
        column_config={"Status": st.column_config.SelectboxColumn(
            options=STATUS_OPTIONS, required=True)})
    if st.button("💾 Simpan kategori ini", key=f"chk_save_{debtor_id}_{cat}", type="primary"):
        conn = get_connection()
        try:
            for _, r in edited.iterrows():
                sub = r["Item"]
                exists = (cat, sub) in docs
                has_data = r["Status"] != "Tidak Ada" or r["Perihal"] or r["Keterangan"]
                if exists or has_data:
                    status = r["Status"]
                    if r["File"] > 0 and status == "Tidak Ada":
                        status = "Ada - Copy/Scan"
                    upsert_doc(conn, debtor_id, cat, sub, status,
                               r["Perihal"] or "", r["Keterangan"] or "")
            conn.commit()
            st.session_state["flash"] = f"Kategori '{cat}' disimpan."
        except Exception as e:
            conn.rollback()
            st.error(f"Gagal menyimpan: {e}")
            conn.close()
            return
        conn.close()
        st.rerun()

    report_section(debtor_id, info, reks, docs)


def tab_profile(debtor_id, info):
    st.subheader("Profil")
    with st.form(f"edit_profile_{debtor_id}"):
        e_nama = st.text_input("Nama Debitur", value=info["nama_debitur"])
        e_rek = st.text_input("Nomor Rekening Utama", value=info["nomor_rekening"] or "")
        e_plafond = st.number_input("Plafond (Rp)", min_value=0.0, step=1_000_000.0,
                                    format="%.0f", value=float(info["plafond"] or 0))
        c1, c2 = st.columns(2)
        e_rm = c1.text_input("RM Kredit", value=info["rm_kredit"] or "")
        e_brimen = c2.text_input("BRIMEN", value=info["brimen"] or "")
        ok = st.form_submit_button("Update Profil", type="primary")
    if ok:
        if not e_nama.strip() or not e_rek.strip():
            st.error("Nama dan rekening utama wajib diisi.")
        else:
            owner = find_duplicate_rek(e_rek.strip(), exclude_id=debtor_id)
            if owner:
                st.error(f"Rekening {e_rek} sudah terdaftar atas nama **{owner}**.")
            else:
                conn = get_connection()
                conn.execute(
                    "UPDATE debtors SET nama_debitur=?, nomor_rekening=?, plafond=?, rm_kredit=?, brimen=? "
                    "WHERE id=?", (e_nama.strip(), e_rek.strip(), e_plafond, e_rm.strip(),
                                   e_brimen.strip(), debtor_id))
                conn.commit()
                conn.close()
                st.session_state["flash"] = "Profil diperbarui."
                st.rerun()

    st.subheader("Rekening Tambahan")
    conn = get_connection()
    accs = pd.read_sql("SELECT * FROM debtor_bank_accounts WHERE debtor_id = ?", conn,
                       params=(debtor_id,))
    conn.close()
    if accs.empty:
        st.caption("Belum ada rekening tambahan.")
    for _, a in accs.iterrows():
        c1, c2 = st.columns([4, 1])
        c1.text(f"💳 {a['nama_bank']} - {a['nomor_rekening']} ({a['keterangan'] or 'Rekening Sekunder'})")
        if c2.button("Hapus", key=f"del_acc_{a['id']}"):
            conn = get_connection()
            conn.execute("DELETE FROM debtor_bank_accounts WHERE id = ?", (int(a["id"]),))
            conn.commit()
            conn.close()
            st.session_state["flash"] = "Rekening dihapus."
            st.rerun()

    with st.expander("➕ Tambah rekening"):
        with st.form(f"add_acc_{debtor_id}"):
            n_num = st.text_input("Nomor Rekening")
            n_bank = st.text_input("Nama Bank", value="BRI")
            n_ket = st.text_input("Keterangan")
            if st.form_submit_button("Simpan Rekening"):
                if not n_num.strip():
                    st.error("Nomor rekening wajib diisi.")
                else:
                    owner = find_duplicate_rek(n_num.strip(), exclude_id=debtor_id)
                    if owner:
                        st.error(f"Rekening sudah terdaftar atas nama **{owner}**.")
                    else:
                        conn = get_connection()
                        conn.execute(
                            "INSERT INTO debtor_bank_accounts (debtor_id, nomor_rekening, nama_bank, keterangan) "
                            "VALUES (?, ?, ?, ?)",
                            (debtor_id, n_num.strip(), n_bank.strip(), n_ket.strip()))
                        conn.commit()
                        conn.close()
                        st.session_state["flash"] = "Rekening ditambahkan."
                        st.rerun()

    st.subheader("Zona Bahaya")
    with st.expander("Hapus debitur ini"):
        st.warning("Menghapus debitur akan menghapus semua data dan SEMUA FILE dokumennya secara permanen.")
        confirm = st.checkbox(f"Saya yakin menghapus {info['nama_debitur']}",
                              key=f"confirm_del_{debtor_id}")
        if st.button("Hapus Debitur Permanen", type="primary", key=f"del_debtor_{debtor_id}"):
            if not confirm:
                st.error("Centang konfirmasi terlebih dahulu.")
            else:
                conn = get_connection()
                rows = conn.execute("SELECT file_path FROM documents WHERE debtor_id = ?",
                                    (debtor_id,)).fetchall()
                for (fp,) in rows:
                    for f in parse_file_paths(fp):
                        try:
                            os.remove(f)
                        except OSError:
                            pass
                shutil.rmtree(DOCS_DIR / str(debtor_id), ignore_errors=True)
                for legacy in DOCS_DIR.glob(f"{debtor_id}_*"):  # folder format lama {id}_{rekening}
                    shutil.rmtree(legacy, ignore_errors=True)
                conn.execute("DELETE FROM debtor_bank_accounts WHERE debtor_id = ?", (debtor_id,))
                conn.execute("DELETE FROM documents WHERE debtor_id = ?", (debtor_id,))
                conn.execute("DELETE FROM debtors WHERE id = ?", (debtor_id,))
                conn.commit()
                conn.close()
                st.session_state.pop("cur_debtor", None)
                st.session_state.pop("debtor_picker", None)
                st.session_state["nav_target"] = PAGE_SEARCH
                st.session_state["flash"] = f"Debitur {info['nama_debitur']} dan semua filenya dihapus."
                st.rerun()


def page_debtor():
    conn = get_connection()
    lst = pd.read_sql("SELECT id, nama_debitur, nomor_rekening FROM debtors "
                      "ORDER BY nama_debitur COLLATE NOCASE", conn)
    conn.close()
    if lst.empty:
        st.info("Belum ada debitur.")
        return

    ids = [int(i) for i in lst["id"]]
    labels = {int(r["id"]): f"{r['nama_debitur']} (Rek: {r['nomor_rekening']})" for _, r in lst.iterrows()}
    cur = st.session_state.get("cur_debtor")
    idx = ids.index(cur) if cur in ids else 0
    debtor_id = st.selectbox("Debitur", ids, index=idx, format_func=lambda i: labels[i],
                             key="debtor_picker")
    if debtor_id != st.session_state.get("cur_debtor"):
        st.session_state["cur_debtor"] = debtor_id
        st.session_state["preview_file"] = None

    conn = get_connection()
    info = pd.read_sql("SELECT * FROM debtors WHERE id = ?", conn, params=(debtor_id,)).iloc[0]
    conn.close()
    docs = load_docs(debtor_id)
    done = sum(1 for (c, s), d in docs.items()
               if (c in CHECKLIST_STRUCTURE and s in CHECKLIST_STRUCTURE[c])
               and is_present(d["status"], [f for f in d["files"] if os.path.exists(f)]))

    st.subheader(info["nama_debitur"])
    summary = load_summary()
    reks = summary.loc[summary["id"] == debtor_id, "rekening"].iloc[0]
    st.caption(f"Rekening: {reks} · Plafond: Rp {float(info['plafond'] or 0):,.0f} · "
               f"RM: {info['rm_kredit'] or '-'} · BRIMEN: {info['brimen'] or '-'}")
    st.progress(min(done / TOTAL_ITEMS, 1.0), text=f"Kelengkapan dokumen {done} / {TOTAL_ITEMS}")

    t_docs, t_up, t_chk, t_prof = st.tabs(
        ["📄 Dokumen", "⬆️ Upload", "✅ Checklist", "✏️ Profil & Rekening"])
    with t_docs:
        tab_documents(debtor_id, docs)
    with t_up:
        tab_upload(debtor_id, docs)
    with t_chk:
        tab_checklist(debtor_id, docs, info, reks)
    with t_prof:
        tab_profile(debtor_id, info)


# =====================================================================
# MAIN
# =====================================================================
st.set_page_config(page_title="SIN - Sistem Informasi Nasabah", layout="wide")

# Terapkan navigasi yang diminta SEBELUM widget navigasi dibuat
if "nav_target" in st.session_state:
    st.session_state["nav_menu"] = st.session_state.pop("nav_target")
if "nav_menu" not in st.session_state:
    st.session_state["nav_menu"] = PAGE_SEARCH

st.sidebar.title("SIN")
st.sidebar.caption("Arsip Dokumen Kredit & Monitoring Debitur")
menu = st.sidebar.radio("Navigasi", [PAGE_SEARCH, PAGE_DEBTOR, PAGE_ADD], key="nav_menu")

if "flash" in st.session_state:
    st.toast(st.session_state.pop("flash"), icon="✅")

if menu == PAGE_SEARCH:
    page_search()
elif menu == PAGE_DEBTOR:
    page_debtor()
else:
    page_add()